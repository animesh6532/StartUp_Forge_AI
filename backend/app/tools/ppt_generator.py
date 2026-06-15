"""PowerPoint generator tool utilizing python-pptx"""

from app.core.logger import logger
from typing import Dict, Any
import io

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


class PPTGenerator:
    """Tool for compiling investor-grade PPTX presentations from multi-agent output data"""

    def __init__(self):
        self.logger = logger

    def generate_pitch_deck(self, content: Dict[str, Any], startup_name: str) -> bytes:
        """Compile and assemble structured multi-slide pitch deck presentation PPTX"""
        self.logger.info(f"Assembling pitch deck PPTX for startup: {startup_name}")
        
        prs = Presentation()
        
        # Configure aspect ratio to widescreen (16:9)
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        # Color codes
        primary_color = RGBColor(79, 70, 229)    # Indigo (#4F46E5)
        secondary_color = RGBColor(13, 148, 136) # Teal (#0D9488)
        dark_neutral = RGBColor(17, 24, 39)      # Charcoal
        light_neutral = RGBColor(243, 244, 246)  # White/Grey
        
        slides_data = content.get("pitchdeck", {}).get("slides", [])
        
        if not slides_data:
            # Fallback mock slides list if pitchdeck agent returned nothing
            slides_data = [
                {
                    "slide_number": 1,
                    "title": f"{startup_name} - Pitch Deck",
                    "subtitle": "Autonomous Multi-Agent Venture Plan.",
                    "points": ["Unlocking scalable operations.", "Driven by AI multi-agents.", "Modern venture planning."],
                    "visual_guideline": "Sleek cover layout."
                }
            ]
            
        for i, slide_info in enumerate(slides_data):
            # Select layouts: 0 = Title Slide, 1 = Title + Content, 5 = Title Only (we use Title Only for complete custom control)
            slide_layout = prs.slide_layouts[5] 
            slide = prs.slides.add_slide(slide_layout)
            
            # --- BACKGROUND COLOR ---
            # Set background to dark neutral for Title slide, light for others
            background = slide.background
            fill = background.fill
            fill.solid()
            if i == 0:
                fill.fore_color.rgb = dark_neutral
            else:
                fill.fore_color.rgb = light_neutral
                
            # --- TITLE BOX ---
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.2))
            tf_title = title_box.text_frame
            tf_title.word_wrap = True
            tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
            
            p_title = tf_title.paragraphs[0]
            p_title.text = slide_info.get("title", "")
            p_title.font.name = "Arial"
            p_title.font.size = Pt(36)
            p_title.font.bold = True
            p_title.font.color.rgb = primary_color if i > 0 else RGBColor(255, 255, 255)
            
            # --- SUBTITLE BOX ---
            subtitle_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(0.8))
            tf_sub = subtitle_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = slide_info.get("subtitle", "")
            p_sub.font.name = "Arial"
            p_sub.font.size = Pt(16)
            p_sub.font.color.rgb = secondary_color if i > 0 else RGBColor(200, 200, 200)
            
            # --- POINTS / CONTENT BOX ---
            content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.7), Inches(3.2))
            tf_content = content_box.text_frame
            tf_content.word_wrap = True
            tf_content.margin_left = tf_content.margin_top = tf_content.margin_right = tf_content.margin_bottom = 0
            
            points = slide_info.get("points", [])
            for k, pt in enumerate(points):
                p_item = tf_content.add_paragraph() if k > 0 else tf_content.paragraphs[0]
                p_item.text = f"•  {pt}"
                p_item.space_after = Pt(14)
                p_item.font.name = "Arial"
                p_item.font.size = Pt(18)
                p_item.font.color.rgb = RGBColor(31, 41, 55) if i > 0 else RGBColor(240, 240, 240)
                
            # --- FOOTER / VISUAL DIRECTION ---
            footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6))
            tf_foot = footer_box.text_frame
            tf_foot.word_wrap = True
            tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
            
            p_foot = tf_foot.paragraphs[0]
            p_foot.text = f"Visual Guideline: {slide_info.get('visual_guideline', '')}"
            p_foot.font.name = "Arial"
            p_foot.font.size = Pt(10)
            p_foot.font.italic = True
            p_foot.font.color.rgb = RGBColor(107, 114, 128) if i > 0 else RGBColor(156, 163, 175)

        # Write to byte buffer
        buffer = io.BytesIO()
        prs.save(buffer)
        ppt_bytes = buffer.getvalue()
        buffer.close()
        
        return ppt_bytes
